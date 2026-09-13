"""Read-only full workbook value regression check before canonical replacement."""
import json
from itertools import zip_longest
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
out = root / "outputs/20260909_ggen_advance_creuset_names"
before = openpyxl.load_workbook(out / "before_ggen_advance_translation_master.xlsx", read_only=True)
after = openpyxl.load_workbook(out / "translation_master_updated.xlsx", read_only=True)
data = json.loads((root / "integrated/translation/ggen_advance_translation_merged.json").read_text(encoding="utf8"))
report = json.loads((out / "manifest.json").read_text(encoding="utf8"))
ids = {j["record_id"] for j in report["jobs"]}
fields = {16: "translation_ko", 17: "translation_status", 18: "translation_source", 19: "review_status", 25: "qa_status", 26: "overlay_batch_id", 27: "translator_notes", 29: "translation_payload_sha256", 34: "translation_segments"}
expected = {}
for n, row in enumerate(data["records"], 7):
    if row["record_id"] in ids:
        for col, key in fields.items():
            value = row.get(key, "")
            if isinstance(value, list):
                value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
            expected[n, col] = value
assert before.sheetnames == after.sheetnames
changes = []
for name in before.sheetnames:
    for n, (a, b) in enumerate(zip_longest(before[name].iter_rows(values_only=True), after[name].iter_rows(values_only=True), fillvalue=()), 1):
        for col, (old, new) in enumerate(zip_longest(a, b), 1):
            if name == "TranslationMaster" and (n, col) in expected:
                assert (new if new is not None else "") == expected[n, col], (name, n, col, new, expected[n, col])
            elif old != new:
                assert name == "Summary" and col == 1 and n <= 6 and data["identity"]["translation_overlay_identity_sha256"] in str(new), (name, n, col, old, new)
            if old != new:
                changes.append([name, n, col])
before.close()
after.close()
(out / "workbook_verification.json").write_text(json.dumps({"result": "PASS", "verified_records": len(ids), "changed_cells": changes, "unrelated_cell_values_unchanged": True}, indent=2) + "\n", encoding="utf8")
print(f"PASS {len(ids)} records, {len(changes)} changed cells; all other workbook values unchanged")
